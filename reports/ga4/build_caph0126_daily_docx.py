"""
Build the CAPH0126 daily-view-pattern docx from the CSVs written by
caph0126_daily.py. Clones the previous topline doc so header, footer and
margins carry over unchanged.

    reports/.venv/bin/python reports/ga4/build_caph0126_daily_docx.py
"""
import csv
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from lib import io

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEMPLATE = Path("/mnt/c/Users/User/Downloads/CAPH0126_MT_eDM_Topline_Results_240826.docx")
OUT_NAME = "CAPH0126_eDM_Daily_Views_270826.docx"

PURPLE = RGBColor(0x70, 0x30, 0xA0)
DARK_NAVY = RGBColor(0x0E, 0x28, 0x41)
MID_GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

GA4 = io.OUT / "ga4"

# GA4 video titles are long; keep the doc's episode labels consistent with the
# page tables.
EP_LABEL = {
    "Ep 1": "Ep 1 - Why Sick Day Planning is Important",
    "Ep 2": "Ep 2 - What Goes in a Sick Day Plan",
    "Ep 3": "Ep 3 - Medication Management",
    "Ep 4": "Ep 4 - Dehydration: The Role of ORS",
    "Ep 5": "Ep 5 - Preparing a Sick Day Kit",
}


def ep_label(title):
    return EP_LABEL.get(title[:4], title)


def read_csv(name):
    with (GA4 / name).open() as f:
        return list(csv.reader(f))


def clear_body(doc):
    body = doc.element.body
    for child in [c for c in body if c.tag != qn("w:sectPr")]:
        body.remove(child)


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_table_borders(tbl, color="D0C0E8"):
    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{side}")
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:color"), color)
                borders.append(el)
            tcPr.append(borders)


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = PURPLE
    return p


def add_para(doc, text, size=9, color=None, bold=False, italic=False, after=4, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = color or MID_GREY
    r.bold = bold
    r.italic = italic
    return p


def add_data_table(doc, headers, rows, col_widths=None, font=9, bold_last=False):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        set_cell_bg(cell, "70309F")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Pt(4)
        r = p.add_run(h)
        r.font.size = Pt(font - 1)
        r.font.color.rgb = WHITE
        r.bold = True

    for ri, row_data in enumerate(rows):
        last = bold_last and ri == len(rows) - 1
        bg = "EDE3F7" if last else ("FFFFFF" if ri % 2 == 0 else "F7F0FC")
        for ci, val in enumerate(row_data):
            cell = tbl.rows[ri + 1].cells[ci]
            set_cell_bg(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Pt(4)
            r = p.add_run(str(val))
            r.font.size = Pt(font)
            r.font.color.rgb = DARK_NAVY
            r.bold = last

    add_table_borders(tbl)
    if col_widths:
        for i, w in enumerate(col_widths):
            for cell in tbl.columns[i].cells:
                cell.width = Inches(w)
    doc.add_paragraph()
    return tbl


def bullets(doc, items):
    for b in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(12)
        p.paragraph_format.first_line_indent = Pt(-12)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        rb = p.add_run("•  ")
        rb.font.color.rgb = PURPLE
        rb.font.size = Pt(9)
        rt = p.add_run(b)
        rt.font.size = Pt(9)
        rt.font.color.rgb = MID_GREY


def short(d):
    return f"{int(d[8:10])}/{int(d[5:7])}"


def main():
    pv = read_csv("caph0126_daily_pageviews.csv")
    by_send = read_csv("caph0126_daily_pageviews_by_send.csv")
    vid = read_csv("caph0126_daily_video_events.csv")

    days = pv[0][2:-1]
    day_hdr = [short(d) for d in days]

    doc = Document(TEMPLATE)
    clear_body(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("eDM Performance: Daily View Pattern")
    r.font.size = Pt(18)
    r.font.color.rgb = PURPLE
    r.bold = True

    add_para(doc, "CAPH0126  |  Hydralyte Diabetes Solus eDM  |  Medicine Today + The Medical Republic",
             size=11, color=DARK_NAVY, bold=True, after=2)
    add_para(doc,
             "Sends: Medicine Today, Thu 20 Aug 2026 (approx 10am AEST)  ·  The Medical Republic, "
             "Tue 25 Aug 2026 (approx 11am AEST)",
             size=9, after=1)
    add_para(doc,
             f"Data window: 7 full days, 20–26 Aug 2026  ·  Prepared: {date.today().strftime('%-d %B %Y')}",
             size=9, after=10)

    # --- 1. daily pageviews -------------------------------------------
    add_heading(doc, "Daily Pageviews by Page")
    add_para(doc, "All traffic to each page, regardless of how the visitor arrived.", after=4)
    add_data_table(
        doc,
        headers=["Page"] + day_hdr + ["Total"],
        rows=[[row[0]] + row[2:] for row in pv[1:]],
        col_widths=[1.85] + [0.45] * len(days) + [0.55],
        font=8,
        bold_last=True,
    )

    # --- 2. split by send ---------------------------------------------
    add_heading(doc, "Split by Send")
    add_para(doc,
             "Both eDMs carry the same campaign tag, but they arrive from different publishers, "
             "so the two sends can still be told apart by traffic source.", after=4)

    sends = ["Medicine Today (20 Aug)", "Medical Republic (25 Aug)", "Other / direct"]
    pages = []
    for row in pv[1:-1]:
        pages.append(row[0])
    totals = {(r[0], r[1]): int(r[-1]) for r in by_send[1:]}
    rows = []
    for pg in pages:
        vals = [totals.get((pg, s), 0) for s in sends]
        rows.append([pg] + [str(v) for v in vals] + [str(sum(vals))])
    col_tot = [sum(totals.get((pg, s), 0) for pg in pages) for s in sends]
    rows.append(["TOTAL"] + [str(v) for v in col_tot] + [str(sum(col_tot))])
    add_data_table(
        doc,
        headers=["Page", "Medicine Today", "Medical Republic", "Other / direct", "Total"],
        rows=rows,
        col_widths=[2.1, 1.0, 1.0, 0.95, 0.6],
        font=8,
        bold_last=True,
    )

    daily = {}
    for r in by_send[1:]:
        for i, d in enumerate(days):
            daily.setdefault(r[1], {}).setdefault(d, 0)
            daily[r[1]][d] += int(r[2 + i])
    add_para(doc, "Daily totals across all seven pages, by send:", after=4, before=2)
    add_data_table(
        doc,
        headers=["Send"] + day_hdr + ["Total"],
        rows=[[s] + [str(daily.get(s, {}).get(d, 0)) for d in days]
              + [str(sum(daily.get(s, {}).values()))] for s in sends],
        col_widths=[1.85] + [0.45] * len(days) + [0.55],
        font=8,
    )

    # --- 3. video -------------------------------------------------------
    add_heading(doc, "Daily Video Views")
    add_para(doc,
             "A view is counted when the player actually starts, so these are lower than pageviews: "
             "some visitors open an episode page without pressing play.", after=4)

    starts = [r for r in vid[1:] if r[1] == "video_start" and r[0].startswith("Ep ")]
    completes = {r[0]: r for r in vid[1:] if r[1] == "video_complete"}
    starts.sort(key=lambda r: r[0])
    add_data_table(
        doc,
        headers=["Episode"] + day_hdr + ["Total"],
        rows=[[ep_label(r[0])] + r[2:] for r in starts],
        col_widths=[1.85] + [0.45] * len(days) + [0.55],
        font=8,
    )

    add_para(doc, "Starts against completions across the whole window:", after=4, before=2)
    crows = []
    for r in starts:
        s_tot = int(r[-1])
        c_tot = int(completes.get(r[0], [0] * len(r))[-1]) if r[0] in completes else 0
        pct = f"{round(c_tot / s_tot * 100)}%" if s_tot else "n/a"
        crows.append([ep_label(r[0]), str(s_tot), str(c_tot), pct])
    ts = sum(int(r[-1]) for r in starts)
    tc = sum(int(completes[r[0]][-1]) for r in starts if r[0] in completes)
    crows.append(["TOTAL", str(ts), str(tc), f"{round(tc / ts * 100)}%" if ts else "n/a"])
    add_data_table(
        doc,
        headers=["Episode", "Starts", "Completions", "Completion rate"],
        rows=crows,
        col_widths=[2.6, 0.8, 1.0, 1.2],
        font=8,
        bold_last=True,
    )

    # --- 4. read-out ----------------------------------------------------
    add_heading(doc, "What the Pattern Shows")

    d0, d1 = days[0], days[1]
    mt, tmr = sends[0], sends[1]
    mt_d0, mt_d1 = daily[mt][d0], daily[mt][d1]
    tmr_send = daily[tmr][days[5]]
    tmr_next = daily[tmr][days[6]]
    mt_tot, tmr_tot = sum(daily[mt].values()), sum(daily[tmr].values())
    grand = int(pv[-1][-1])
    cb = next(int(r[-1]) for r in pv[1:] if r[0].startswith("Clinical Bites"))
    ep1 = next(int(r[-1]) for r in pv[1:] if r[0].startswith("Ep 1"))
    ep5 = next(int(r[-1]) for r in pv[1:] if r[0].startswith("Ep 5"))
    low = min(int(pv[-1][2 + i]) for i in range(len(days)))
    low_day = days[[int(pv[-1][2 + i]) for i in range(len(days))].index(low)]
    peak = int(pv[-1][2])

    bullets(doc, [
        f"Both sends spike hard on the send day and fall away fast. Medicine Today put {mt_d0} views "
        f"across the seven pages on day one, then {mt_d1} the next day. The Medical Republic put "
        f"{tmr_send} on its send day and was still running at {tmr_next} the day after.",
        f"The weekend is close to dead: {low} views across all seven pages on "
        f"{date.fromisoformat(low_day).strftime('%A %-d %b')}, against {peak} on the Thursday send day.",
        f"The Clinical Bites landing page is the workhorse, taking {cb} of the {grand} views. It is "
        "where the Watch Now button lands, and it out-pulls every individual episode.",
        f"Episode 1 takes {ep1} views, Episode 5 takes {ep5}. The drop is steepest between Episode 1 "
        "and Episode 2, then it flattens rather than continuing to fall.",
        f"Medicine Today has delivered {mt_tot} views against the Medical Republic's {tmr_tot}, but "
        "the Medical Republic has had two days to run against Medicine Today's seven.",
        "Episode 3 slightly out-pulls Episode 2 on both views and starts, so viewers are not strictly "
        "working through the series in order.",
    ])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r1 = p.add_run("Data source: ")
    r1.font.size = Pt(8)
    r1.font.color.rgb = MID_GREY
    r1.bold = True
    r2 = p.add_run(
        "GA4 property 306115293 (hcp.carepharma.com.au). Pageviews are screen/page views, video figures "
        "are the player's own start and complete events. The window closes at the end of Wednesday 26 Aug; "
        "GA4 had not yet processed Thursday 27 Aug at the time of writing. Both sends share "
        "utm_campaign=CAPH0126, so the split is taken from traffic source, not campaign tag."
    )
    r2.font.size = Pt(8)
    r2.font.color.rgb = MID_GREY

    out = io.OUT / OUT_NAME
    doc.save(out)
    print(f"wrote {out.relative_to(io.PROJECT_ROOT)}")
    io.copy_to_downloads(out)


if __name__ == "__main__":
    main()
