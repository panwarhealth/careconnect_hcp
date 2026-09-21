"""
Build the combined CAPH0126 topline docx: Medicine Today alongside
The Medical Republic.

Medicine Today figures are reproduced verbatim from the 24 Aug topline that
was already circulated, so this doc cannot disagree with it. The Medical
Republic figures are pulled through the identical GA4 queries over its own
first five days, which keeps the two columns comparable.

    reports/.venv/bin/python reports/ga4/build_caph0126_combined_docx.py
"""
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from lib import io

from build_caph0126_daily_docx import (
    DARK_NAVY,
    MID_GREY,
    PURPLE,
    add_data_table,
    add_heading,
    add_para,
    bullets,
    clear_body,
)

from docx import Document
from docx.shared import Pt

TEMPLATE = Path("/mnt/c/Users/User/Downloads/CAPH0126_MT_eDM_Topline_Results_240826.docx")
OUT_NAME = f"CAPH0126_eDM_Topline_Results_MT_vs_TMR_{date.today():%d%m%y}.docx"

# --- Traffic & Reach: label, MT, TMR --------------------------------------
# MT column verbatim from the 24 Aug topline; TMR from caph0126_by_send.py.
REACH = [
    ("Site visits (5 days post-send)", "278", "276"),
    ("Visits from the email", "202", "164"),
    ("People reached (unique users)", "184", "147"),
    ("First-time visitors to the site", "195", "152"),
    ("Engaged visits", "164 (81%)", "139 (85%)"),
    ("Avg time actively reading", "~1.0 min", "~1.1 min"),
    ("New HCP registrations", "11", "12"),
    ("Video starts (Clinical Bites series)", "105", "48"),
]

# CTA rows: label, MT visits, TMR visits. Percentages are each send's share
# of all link-driven visits (MT 363, TMR 291 incl. logo/header links).
CTAS = [
    ("Watch Now (videos)", 146, 92),
    ("Download Sick Day Plan", 100, 76),
    ("Take the Quiz", 34, 28),
    ("Read More: anti-obesity medications article", 15, 29),
    ("Download Patient Leaflet", 20, 6),
    ("Order Samples", 17, 16),
    ("Read More: POTS article", 16, 16),
]
MT_LINK_TOTAL, TMR_LINK_TOTAL = 363, 291

# Episode starts: label, MT, TMR.
EPISODES = [
    ("Ep 1: Why Sick Day Planning is Important", 45, 18),
    ("Ep 2: What Goes in a Sick Day Plan", 18, 8),
    ("Ep 3: Medication Management During Sick Days", 22, 9),
    ("Ep 4: Dehydration and the Role of ORS", 12, 6),
    ("Ep 5: Preparing a Sick Day Management Kit", 8, 7),
]


def pct(n, total):
    return f"{n} ({round(n / total * 100)}%)"


def main() -> None:
    doc = Document(TEMPLATE)
    clear_body(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("eDM Results: Medicine Today vs The Medical Republic")
    r.font.size = Pt(18)
    r.font.color.rgb = PURPLE
    r.bold = True

    add_para(doc, "CAPH0126  |  Hydralyte Diabetes Solus eDM",
             size=11, color=DARK_NAVY, bold=True, after=2)
    add_para(doc,
             "Sends: Medicine Today, Thu 20 Aug 2026  ·  The Medical Republic, Tue 25 Aug 2026  ·  "
             "Both emails measured over their first 5 days",
             size=9, after=1)
    add_para(doc, "Prepared: 31 August 2026", size=9, after=10)

    add_heading(doc, "Traffic & Reach")
    add_para(doc,
             "For context: a normal week on the site before either send ran at 136 visits across "
             "seven days, so both emails lifted the site to roughly double its usual weekly "
             "traffic in five days.", after=4)
    add_data_table(
        doc,
        headers=["", "Medicine Today", "The Medical Republic"],
        rows=[list(row) for row in REACH],
        col_widths=[2.9, 1.5, 1.7],
        font=9,
    )
    add_para(doc,
             "Medicine Today figures are as reported on 24 August. The Medical Republic is "
             "measured the same way over its own five days (25 to 29 Aug).",
             after=6, italic=True)

    add_heading(doc, "Which Links People Clicked")
    add_para(doc,
             "Counted as visits to the site arriving via each button, not clicks recorded by the "
             "publisher. Percentages are each email's share of all link-driven visits.", after=4)
    add_data_table(
        doc,
        headers=["Link in the email", "Medicine Today", "The Medical Republic"],
        rows=[[label, pct(mt, MT_LINK_TOTAL), pct(tmr, TMR_LINK_TOTAL)]
              for label, mt, tmr in CTAS],
        col_widths=[2.9, 1.5, 1.7],
        font=9,
    )
    add_para(doc,
             "The Sick Day Plan download drew the most first-time visitors in both sends, even "
             "though Watch Now drew more visits overall.",
             after=6, italic=True)

    add_heading(doc, "Video Engagement (Clinical Bites)")
    add_para(doc,
             "Five-episode series at /clinical-bites/, ungated across the send windows. "
             "A start is counted when the player actually plays.", after=4)
    add_data_table(
        doc,
        headers=["Episode", "MT starts", "TMR starts"],
        rows=[[label, str(mt), str(tmr)] for label, mt, tmr in EPISODES]
        + [["TOTAL", "105", "48"]],
        col_widths=[3.4, 1.3, 1.4],
        font=9,
        bold_last=True,
    )
    add_para(doc,
             "Around 1 in 4 viewers who started an episode watched it to the end, similar for "
             "both sends.",
             after=6, italic=True)

    add_heading(doc, "What It Means")
    bullets(doc, [
        "Registrations came out level: 11 new HCPs from Medicine Today, 12 from The Medical "
        "Republic, and in both cases almost all arrived on the send day itself.",
        "The Medical Republic delivered about 80% of Medicine Today's traffic, with slightly "
        "higher engagement (85% engaged visits vs 81%). Its traffic faded faster though: by day "
        "five it was down to 1 visit, where Medicine Today still had 13.",
        "Watch Now was the top link in both emails, but the audiences diverge below that: "
        "The Medical Republic readers went for the anti-obesity medications article (29 visits "
        "vs 15) and largely ignored the patient leaflet (6 vs 20).",
        "Video interest was thinner from The Medical Republic: 48 starts against 105, a bigger "
        "gap than overall traffic explains.",
    ])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r1 = p.add_run("Data sources: ")
    r1.font.size = Pt(8)
    r1.font.color.rgb = MID_GREY
    r1.bold = True
    r2 = p.add_run(
        "GA4 property 306115293 (visits, engagement, link attribution, video events) and the "
        "site registration database; registration counts exclude internal accounts. Both sends "
        "share the same campaign tag and are separated by traffic source. Link-visit totals "
        "exceed each email's overall visit total because a visit using more than one link is "
        "counted under each."
    )
    r2.font.size = Pt(8)
    r2.font.color.rgb = MID_GREY

    out = io.OUT / OUT_NAME
    doc.save(out)
    print(f"wrote {out.relative_to(io.PROJECT_ROOT)}")
    io.copy_to_downloads(out)


if __name__ == "__main__":
    main()
